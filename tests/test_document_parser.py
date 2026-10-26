import pytest
from backend.app.services.document_parser import DocumentParser, ParsedDocument

@pytest.fixture
def parser():
    return DocumentParser()

def test_parse_txt(parser):
    content = b"DocuMind AI is an enterprise-grade document intelligence platform.\nIt uses vector search for RAG."
    doc = parser.parse("overview.txt", content)
    assert isinstance(doc, ParsedDocument)
    assert doc.filename == "overview.txt"
    assert doc.file_type == "txt"
    assert doc.total_pages == 1
    assert "DocuMind AI" in doc.pages[0].text

def test_parse_empty_unsupported_file(parser):
    with pytest.raises(ValueError, match="Unsupported file format"):
        parser.parse("unsupported.xyz", b"hello world")

def test_parse_txt_encodings(parser):
    # Test latin-1 encoded text with special characters
    content = "DocuMind café and résumé testing".encode("latin-1")
    doc = parser.parse("resume.txt", content)
    assert "caf" in doc.pages[0].text

def test_parse_docx_mock_or_real(parser):
    import io
    try:
        import docx
        doc_obj = docx.Document()
        doc_obj.add_paragraph("This is paragraph 1 of the testing document.")
        table = doc_obj.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "Header 1"
        table.cell(0, 1).text = "Header 2"
        table.cell(1, 0).text = "Val 1"
        table.cell(1, 1).text = "Val 2"
        buf = io.BytesIO()
        doc_obj.save(buf)
        content = buf.getvalue()

        parsed = parser.parse("sample.docx", content)
        assert parsed.file_type == "docx"
        assert "paragraph 1" in parsed.pages[0].text
        assert "Header 1" in parsed.pages[0].text
    except ImportError:
        pytest.skip("docx library not installed")

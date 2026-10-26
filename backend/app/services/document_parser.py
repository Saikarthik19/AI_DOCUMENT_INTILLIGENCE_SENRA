import io
import os
from typing import List, Dict, Any, Optional
from dataclasses import dataclass
from backend.app.core.logging import logger
from backend.app.services.ocr_service import ocr_service

# Optional imports with fallbacks
try:
    import fitz  # PyMuPDF
    PYMUPDF_AVAILABLE = True
except ImportError:
    PYMUPDF_AVAILABLE = False

try:
    import pypdf
    PYPDF_AVAILABLE = True
except ImportError:
    PYPDF_AVAILABLE = False

try:
    import docx
    DOCX_AVAILABLE = True
except ImportError:
    DOCX_AVAILABLE = False

@dataclass
class ParsedPage:
    page_number: int
    text: str
    metadata: Dict[str, Any]

@dataclass
class ParsedDocument:
    filename: str
    file_type: str
    total_pages: int
    pages: List[ParsedPage]
    metadata: Dict[str, Any]

class DocumentParser:
    def __init__(self):
        logger.info("Initialized DocumentParser (PyMuPDF: %s, PyPDF: %s, DOCX: %s)",
                    PYMUPDF_AVAILABLE, PYPDF_AVAILABLE, DOCX_AVAILABLE)

    def parse(self, filename: str, content: bytes, file_type: Optional[str] = None) -> ParsedDocument:
        ext = os.path.splitext(filename)[1].lower()
        if not file_type:
            file_type = ext.replace(".", "")

        if ext == ".pdf":
            return self._parse_pdf(filename, content)
        elif ext == ".docx":
            return self._parse_docx(filename, content)
        elif ext == ".txt":
            return self._parse_txt(filename, content)
        elif ext in [".png", ".jpg", ".jpeg"]:
            return self._parse_image(filename, content)
        else:
            raise ValueError(f"Unsupported file format: '{ext}'. Allowed: .pdf, .docx, .txt, .png, .jpg, .jpeg")

    def _parse_pdf(self, filename: str, content: bytes) -> ParsedDocument:
        pages: List[ParsedPage] = []
        
        # Primary: PyMuPDF
        if PYMUPDF_AVAILABLE:
            try:
                doc = fitz.open(stream=content, filetype="pdf")
                if doc.is_encrypted:
                    raise ValueError(f"PDF document '{filename}' is encrypted or password-protected.")
                
                total_pages = len(doc)
                if total_pages == 0:
                    raise ValueError(f"PDF document '{filename}' contains 0 pages.")

                for page_idx in range(total_pages):
                    page = doc[page_idx]
                    page_num = page_idx + 1
                    text = page.get_text("text").strip()
                    
                    # Detect low-text or scanned pages
                    is_scanned = False
                    if len(text) < 30 and ocr_service.is_available():
                        # Render page pixmap for OCR
                        try:
                            pix = page.get_pixmap(dpi=150)
                            img_bytes = pix.tobytes("png")
                            ocr_text = ocr_service.extract_text_from_image_bytes(img_bytes)
                            if len(ocr_text) > len(text):
                                text = ocr_text
                                is_scanned = True
                        except Exception as e:
                            logger.debug("Page %s OCR fallback error: %s", page_num, str(e))

                    pages.append(ParsedPage(
                        page_number=page_num,
                        text=text,
                        metadata={"is_scanned": is_scanned, "char_count": len(text)}
                    ))
                doc.close()
                return ParsedDocument(
                    filename=filename,
                    file_type="pdf",
                    total_pages=total_pages,
                    pages=pages,
                    metadata={"parser": "PyMuPDF"}
                )
            except Exception as e:
                logger.warning("PyMuPDF parsing failed (%s), attempting fallback if possible...", str(e))
                if not PYPDF_AVAILABLE:
                    raise ValueError(f"Failed to parse PDF '{filename}': {str(e)}")

        # Fallback: pypdf
        if PYPDF_AVAILABLE:
            try:
                reader = pypdf.PdfReader(io.BytesIO(content))
                if reader.is_encrypted:
                    raise ValueError(f"PDF document '{filename}' is encrypted.")
                total_pages = len(reader.pages)
                for idx, p in enumerate(reader.pages):
                    extracted = p.extract_text() or ""
                    pages.append(ParsedPage(
                        page_number=idx + 1,
                        text=extracted.strip(),
                        metadata={"is_scanned": False, "char_count": len(extracted)}
                    ))
                return ParsedDocument(
                    filename=filename,
                    file_type="pdf",
                    total_pages=total_pages,
                    pages=pages,
                    metadata={"parser": "pypdf"}
                )
            except Exception as e:
                raise ValueError(f"Could not parse PDF '{filename}': {str(e)}")

        raise ValueError("No PDF parser library (PyMuPDF or pypdf) is available.")

    def _parse_docx(self, filename: str, content: bytes) -> ParsedDocument:
        if not DOCX_AVAILABLE:
            raise ValueError("python-docx is not installed to parse .docx files.")
        try:
            doc = docx.Document(io.BytesIO(content))
            full_text_parts = []
            
            # Paragraphs
            for p in doc.paragraphs:
                p_text = p.text.strip()
                if p_text:
                    full_text_parts.append(p_text)

            # Tables
            for table in doc.tables:
                table_lines = []
                for row in table.rows:
                    row_cells = [cell.text.strip() for cell in row.cells]
                    if any(row_cells):
                        table_lines.append(" | ".join(row_cells))
                if table_lines:
                    full_text_parts.append("\n[TABLE]\n" + "\n".join(table_lines) + "\n[/TABLE]\n")

            combined_text = "\n\n".join(full_text_parts).strip()
            pages = [ParsedPage(page_number=1, text=combined_text, metadata={"paragraphs": len(doc.paragraphs), "tables": len(doc.tables)})]
            return ParsedDocument(
                filename=filename,
                file_type="docx",
                total_pages=1,
                pages=pages,
                metadata={"tables_count": len(doc.tables)}
            )
        except Exception as e:
            raise ValueError(f"Failed to read DOCX file '{filename}': {str(e)}")

    def _parse_txt(self, filename: str, content: bytes) -> ParsedDocument:
        # Multi-encoding decode
        encodings = ["utf-8", "utf-8-sig", "latin-1", "cp1252", "iso-8859-1"]
        decoded_text = None
        for enc in encodings:
            try:
                decoded_text = content.decode(enc)
                break
            except UnicodeDecodeError:
                continue

        if decoded_text is None:
            decoded_text = content.decode("utf-8", errors="replace")

        cleaned = decoded_text.strip()
        pages = [ParsedPage(page_number=1, text=cleaned, metadata={"char_count": len(cleaned)})]
        return ParsedDocument(
            filename=filename,
            file_type="txt",
            total_pages=1,
            pages=pages,
            metadata={"encoding": "auto"}
        )

    def _parse_image(self, filename: str, content: bytes) -> ParsedDocument:
        if not ocr_service.is_available():
            raise ValueError(
                f"Cannot process image '{filename}': Tesseract OCR is not installed or not available on the server. "
                "Please configure Tesseract OCR to enable image extraction."
            )
        text = ocr_service.extract_text_from_image_bytes(content)
        pages = [ParsedPage(page_number=1, text=text, metadata={"ocr": True, "char_count": len(text)})]
        return ParsedDocument(
            filename=filename,
            file_type="image",
            total_pages=1,
            pages=pages,
            metadata={"ocr": True}
        )

document_parser = DocumentParser()
